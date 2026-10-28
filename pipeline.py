"""Agentic pipeline: Planner -> Generator (fine-tuned LLM) -> Critic -> (retry loop) -> Formatter."""
import random, re
from dataclasses import dataclass, field
from common import SYSTEM, make_user_prompt, parse_user_prompt, split_output
from readability import fk_grade

UNSAFE = {"kill", "killed", "blood", "gun", "die", "died", "death", "dead", "hate", "stupid", "idiot",
          "sex", "drug", "drugs", "alcohol", "beer", "murder", "weapon", "suicide", "bomb"}
REQUIRED_SECTIONS = ["Objectives", "Materials", "Activities", "Discussion Questions", "Assessment"]


# ---------------------------------------------------------------- generator backends
class TemplateBackend:
    """Offline stand-in so the whole app/pipeline runs with no GPU or model download.
    It re-uses the training-data renderer. Use HFBackend for the real fine-tuned model."""
    name = "template-demo"

    def __init__(self):
        self.calls = 0

    def generate(self, messages, temperature=0.7, max_new_tokens=900):
        from curriculum import TOPICS
        from build_dataset import render, NAMES, PLACES, GUIDES
        req = parse_user_prompt(messages[1]["content"])
        kws = req["keywords"] or [req["topic"] or "learning"]
        best, score = None, 0
        for t in TOPICS:
            sc = len({k.lower() for k in kws} & {k.lower() for k in t["keywords"]})
            sc += 3 * (t["topic"].lower() == req["topic"].lower())
            if sc > score:
                best, score = t, sc
        if best is None:  # unseen topic -> generic content
            tp = req["topic"] or kws[0]
            best = dict(grade=req["grade"], subject=req["subject"] or "General", topic=tp,
                        objectives=[f"Explain the main idea of {tp}", f"Use the words {', '.join(kws)} correctly"],
                        facts=[f"{tp} is something we can observe and talk about.",
                               f"Looking closely at {kws[0]} helps us understand it."],
                        activity=f"Observe and discuss {tp} with a partner and draw what you see",
                        materials=["notebook", "colour pencils"],
                        questions=[f"What did you learn about {tp}?"])
        self.calls += 1
        rng = random.Random(self.calls * 7919 + int(temperature * 100))
        return render(best, rng.choice(NAMES), rng.choice(PLACES), rng.choice(GUIDES), kws, req["grade"], rng)


class HFBackend:
    """Base model + optional LoRA adapter produced by train_qlora.py."""
    def __init__(self, base="Qwen/Qwen2.5-3B-Instruct", adapter=None):
        import torch
        from transformers import AutoTokenizer, AutoModelForCausalLM, BitsAndBytesConfig
        self.name = f"hf:{base}" + (f"+{adapter}" if adapter else " (base)")
        self.tok = AutoTokenizer.from_pretrained(base)
        kw = {}
        if torch.cuda.is_available():
            kw = dict(device_map="auto", quantization_config=BitsAndBytesConfig(
                load_in_4bit=True, bnb_4bit_quant_type="nf4", bnb_4bit_compute_dtype=torch.float16))
        self.model = AutoModelForCausalLM.from_pretrained(base, **kw)
        if adapter:
            from peft import PeftModel
            self.model = PeftModel.from_pretrained(self.model, adapter)
        self.model.eval()

    def generate(self, messages, temperature=0.7, max_new_tokens=900):
        import torch
        prompt = self.tok.apply_chat_template(messages, tokenize=False, add_generation_prompt=True)
        ids = self.tok(prompt, return_tensors="pt").to(self.model.device)
        with torch.no_grad():
            out = self.model.generate(**ids, max_new_tokens=max_new_tokens, do_sample=True,
                                      temperature=max(temperature, 0.1), top_p=0.9,
                                      repetition_penalty=1.05, pad_token_id=self.tok.eos_token_id)
        return self.tok.decode(out[0][ids["input_ids"].shape[1]:], skip_special_tokens=True).strip()


# ---------------------------------------------------------------- planner
@dataclass
class Plan:
    grade: int
    subject: str
    topic: str
    keywords: list
    max_fk: float
    messages: list = field(default_factory=list)
    dropped: list = field(default_factory=list)   # requested keywords removed by the input safety guard


def planner(grade, subject, topic, keywords):
    keywords = [k.strip() for k in (keywords.split(",") if isinstance(keywords, str) else keywords) if k.strip()]
    # Input safety guard: never pass unsuitable keywords to the model (also checked again by the critic)
    dropped = [k for k in keywords if set(re.findall(r"[a-z']+", k.lower())) & UNSAFE]
    keywords = [k for k in keywords if k not in dropped]
    user = make_user_prompt(grade, subject, topic, keywords)
    # FK is inflated by subject terms (e.g. "evaporation"), so the limit is grade + 4 (calibrated on the dataset)
    return Plan(int(grade), subject, topic, keywords, max_fk=int(grade) + 4.0, dropped=dropped,
                messages=[{"role": "system", "content": SYSTEM}, {"role": "user", "content": user}])


# ---------------------------------------------------------------- critic
def critic(plan, raw):
    story, lesson = split_output(raw)
    issues = []
    if not story or not lesson:
        issues.append("Output must contain both 'STORY:' and 'LESSON PLAN:' sections.")
    for sec in REQUIRED_SECTIONS:
        if lesson and sec.lower() not in lesson.lower():
            issues.append(f"Lesson plan is missing the '{sec}' section.")
    low = raw.lower()
    missing = [k for k in plan.keywords if k.lower() not in low]
    if missing:
        issues.append("Keywords not used: " + ", ".join(missing))
    words = re.findall(r"[a-z']+", low)
    bad = sorted(set(words) & UNSAFE)
    if bad:
        issues.append("Unsuitable words for children: " + ", ".join(bad))
    fk = fk_grade(story) if story else 0.0
    if story and fk > plan.max_fk:
        issues.append(f"Story reading level (grade {fk:.1f}) is too high for grade {plan.grade}; use shorter sentences and simpler words.")
    n = len(story.split())
    if story and not (60 <= n <= 450):
        issues.append(f"Story length is {n} words; keep it between 60 and 450.")
    metrics = {"fk_grade": round(fk, 2), "story_words": n, "keywords_covered": f"{len(plan.keywords) - len(missing)}/{len(plan.keywords)}"}
    return issues, metrics


# ---------------------------------------------------------------- formatter
def to_markdown(plan, story, lesson):
    return (f"# {plan.topic} (Grade {plan.grade}, {plan.subject})\n\n## Story\n\n{story}\n\n"
            f"## Lesson Plan\n\n{lesson}\n").replace("\nObjectives:", "\n**Objectives:**")


def to_docx(plan, story, lesson, path):
    try:
        from docx import Document
    except ImportError:
        return None
    d = Document()
    d.add_heading(f"{plan.topic} - Grade {plan.grade} {plan.subject}", 0)
    d.add_heading("Story", 1); d.add_paragraph(story)
    d.add_heading("Lesson Plan", 1)
    for line in lesson.splitlines():
        d.add_paragraph(line)
    d.save(path)
    return path


# ---------------------------------------------------------------- orchestration
def run(backend, grade, subject, topic, keywords, max_retries=2, agentic=True):
    plan = planner(grade, subject, topic, keywords)
    messages = list(plan.messages)
    trace, raw, issues, metrics = [], "", [], {}
    for attempt in range(1 + (max_retries if agentic else 0)):
        raw = backend.generate(messages, temperature=0.6 + 0.15 * attempt)
        issues, metrics = critic(plan, raw)
        trace.append({"attempt": attempt + 1, "issues": issues, **metrics})
        if not issues or not agentic:
            break
        messages = list(plan.messages) + [
            {"role": "assistant", "content": raw},
            {"role": "user", "content": "Rewrite it fixing these problems:\n- " + "\n- ".join(issues)}]
    story, lesson = split_output(raw)
    if plan.dropped:
        metrics = {**metrics, "keywords_removed_for_safety": plan.dropped}
    return dict(plan=plan, story=story, lesson=lesson, raw=raw, issues=issues,
                metrics=metrics, trace=trace, passed=not issues, dropped=plan.dropped)
